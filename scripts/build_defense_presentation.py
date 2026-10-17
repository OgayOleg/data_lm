#!/usr/bin/env python
"""Build the 10-slide PowerPoint deck for the 5–7 minute project defense."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "reports" / "figures"
OUTPUT = ROOT / "reports" / "defense_presentation.pptx"

NAVY = RGBColor(15, 35, 63)
BLUE = RGBColor(43, 108, 176)
CYAN = RGBColor(38, 166, 197)
GREEN = RGBColor(47, 133, 90)
ORANGE = RGBColor(192, 86, 33)
RED = RGBColor(197, 48, 48)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(247, 249, 252)
PALE_BLUE = RGBColor(238, 244, 252)
PALE_GREEN = RGBColor(238, 247, 241)
MUTED = RGBColor(100, 116, 139)
GRID = RGBColor(222, 230, 240)
DARK_BG = RGBColor(8, 20, 38)

SLIDE_W = 13.333
SLIDE_H = 7.5


def set_text(tf, text, size=18, color=NAVY, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tf


def add_text(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(shape.text_frame, text, size, color, bold, align=align)
    return shape


def add_box(slide, x, y, w, h, fill=LIGHT, line=GRID, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_panel(slide, x, y, w, h, title, body, accent=BLUE, body_size=15):
    box = add_box(slide, x, y, w, h)
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, x + 0.22, y + 0.14, w - 0.4, 0.34, title, 16, NAVY, True)
    add_text(slide, x + 0.22, y + 0.58, w - 0.4, h - 0.72, body, body_size, MUTED)
    return box


def add_metric(slide, x, y, w, label, value, accent=BLUE, note=""):
    add_box(slide, x, y, w, 1.03, PALE_BLUE if accent == BLUE else PALE_GREEN, GRID)
    add_text(slide, x + 0.12, y + 0.11, w - 0.24, 0.25, label.upper(), 10, MUTED, True)
    add_text(slide, x + 0.12, y + 0.34, w - 0.24, 0.39, value, 24, accent, True)
    if note:
        add_text(slide, x + 0.12, y + 0.75, w - 0.24, 0.18, note, 9, MUTED)


def add_title(slide, number, title, subtitle=""):
    add_text(slide, 0.55, 0.25, 0.56, 0.42, f"{number:02d}", 16, BLUE, True)
    add_text(slide, 1.15, 0.22, 11.55, 0.5, title, 25, NAVY, True)
    if subtitle:
        add_text(slide, 1.16, 0.68, 11.2, 0.3, subtitle, 11, MUTED)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.03), Inches(12.2), Inches(0.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = GRID
    line.line.fill.background()


def add_footer(slide, index, role="Совместно"):
    add_text(slide, 0.55, 7.16, 5.2, 0.18, f"Введение в науку о данных  •  {role}", 9, MUTED)
    add_text(slide, 12.15, 7.16, 0.6, 0.18, f"{index}/10", 9, MUTED, True, PP_ALIGN.RIGHT)


def add_bullets(slide, x, y, w, h, items, size=15, color=NAVY, spacing=7):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(5)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.text = "•  " + p.text
    return shape


def add_picture_fit(slide, filename, x, y, w, h):
    path = FIG / filename
    with Image.open(path) as image:
        ratio = image.width / image.height
    area_ratio = w / h
    if ratio > area_ratio:
        pic_w = w
        pic_h = w / ratio
        pic_x, pic_y = x, y + (h - pic_h) / 2
    else:
        pic_h = h
        pic_w = h * ratio
        pic_x, pic_y = x + (w - pic_w) / 2, y
    return slide.shapes.add_picture(
        str(path), Inches(pic_x), Inches(pic_y), width=Inches(pic_w), height=Inches(pic_h)
    )


def style_table(table, header_fill=NAVY, font_size=11, first_col_wide=False):
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.color.rgb = WHITE if r == 0 else NAVY
                p.font.bold = r == 0
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER


def add_table(slide, x, y, w, h, data, widths=None, font_size=11):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    style_table(table, font_size=font_size)
    return table


def add_flow(slide, labels, x=0.75, y=2.28, total_w=11.85):
    gap = 0.25
    box_w = (total_w - gap * (len(labels) - 1)) / len(labels)
    for i, (label, accent) in enumerate(labels):
        bx = x + i * (box_w + gap)
        add_box(slide, bx, y, box_w, 1.0, WHITE, accent)
        add_text(slide, bx + 0.08, y + 0.25, box_w - 0.16, 0.5, label, 13, NAVY, True, PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            add_text(slide, bx + box_w, y + 0.31, gap, 0.35, "→", 17, BLUE, True, PP_ALIGN.CENTER)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # 1 — title
    slide = blank_slide(prs)
    slide.background.fill.fore_color.rgb = DARK_BG
    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), Inches(SLIDE_H)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CYAN
    accent.line.fill.background()
    add_text(slide, 0.8, 0.72, 5.8, 0.38, "ПРОЕКТ-ПРАКТИКУМ  •  DEVOPS-ИНЖЕНЕРИЯ", 12, CYAN, True)
    add_text(
        slide,
        0.8,
        1.28,
        11.65,
        1.55,
        "Разработка и автоматизация\nML-пайплайна бинарной классификации",
        34,
        WHITE,
        True,
    )
    add_text(
        slide,
        0.82,
        3.03,
        10.8,
        0.48,
        "Telco Customer Churn  ×  Bank Marketing",
        20,
        RGBColor(187, 210, 238),
        True,
    )
    add_panel(
        slide,
        0.82,
        4.05,
        5.65,
        1.34,
        "Data Scientist / Researcher",
        "ФИО: ____________________\nEDA · признаки · модели · интерпретация",
        CYAN,
        14,
    )
    add_panel(
        slide,
        6.78,
        4.05,
        5.65,
        1.34,
        "ML Engineer / DevOps Engineer",
        "ФИО: ____________________\nархитектура · тесты · DVC · CI · документация",
        GREEN,
        14,
    )
    add_text(
        slide, 0.82, 6.32, 7.6, 0.34, "Дисциплина: «Введение в науку о данных»", 13, RGBColor(187, 210, 238)
    )
    add_text(
        slide, 9.6, 6.32, 2.82, 0.34, "Группа: __________", 13, RGBColor(187, 210, 238), False, PP_ALIGN.RIGHT
    )
    add_text(slide, 11.86, 7.16, 0.6, 0.18, "1/10", 9, RGBColor(187, 210, 238), True, PP_ALIGN.RIGHT)

    # 2 — task and data
    slide = blank_slide(prs)
    add_title(slide, 2, "Постановка задачи и данные", "Два публичных датасета Kaggle, 48 231 наблюдение")
    add_footer(slide, 2, "Data Scientist")
    add_panel(
        slide,
        0.65,
        1.3,
        5.95,
        2.28,
        "TELCO CUSTOMER CHURN",
        (
            "Бизнес-задача: кто из клиентов уйдёт?\n\n"
            "7 043 записи  •  20 признаков\nЦель: Churn = Yes / No\n"
            "Положительный класс: 26,5%\nОсобенность: 11 пропусков TotalCharges"
        ),
        BLUE,
        15,
    )
    add_panel(
        slide,
        6.75,
        1.3,
        5.95,
        2.28,
        "BANK MARKETING",
        (
            "Бизнес-задача: кто подпишется на депозит?\n\n"
            "41 188 записей  •  20 признаков\nЦель: y = yes / no\n"
            "Положительный класс: 11,3%\nОсобенность: утечка в duration"
        ),
        GREEN,
        15,
    )
    add_picture_fit(slide, "telco_target_balance.png", 0.8, 3.88, 5.55, 2.62)
    add_picture_fit(slide, "bank_target_balance.png", 6.95, 3.88, 5.55, 2.62)
    add_text(
        slide,
        0.8,
        6.58,
        11.65,
        0.34,
        "Почему два датасета: разные уровни дисбаланса и разные риски качества данных.",
        14,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 3 — EDA
    slide = blank_slide(prs)
    add_title(
        slide,
        3,
        "EDA: ключевые находки",
        "От описания данных — к проверяемым гипотезам и клиентским сегментам",
    )
    add_footer(slide, 3, "Data Scientist")
    add_picture_fit(slide, "telco_correlation_heatmap.png", 0.65, 1.28, 6.15, 4.45)
    add_panel(
        slide,
        7.05,
        1.3,
        5.62,
        1.1,
        "42,7% ОТТОКА",
        "у клиентов с помесячным контрактом против 2,8% у Two year",
        ORANGE,
        14,
    )
    add_panel(
        slide,
        7.05,
        2.58,
        5.62,
        1.1,
        "~50% ОТТОКА",
        "у клиентов со стажем 0–10 месяцев; после 4 лет — около 10%",
        BLUE,
        14,
    )
    add_panel(
        slide,
        7.05,
        3.86,
        5.62,
        1.1,
        "45,3% ОТТОКА",
        "при оплате Electronic check против ~15% у автоплатежей",
        GREEN,
        14,
    )
    add_panel(
        slide,
        7.05,
        5.14,
        5.62,
        1.1,
        "ГЛАВНЫЙ СЕГМЕНТ РИСКА",
        "новый клиент + месячный контракт + ручная оплата",
        RED,
        14,
    )
    add_text(
        slide,
        0.8,
        6.5,
        11.65,
        0.34,
        "Практический результат EDA: адресный список для кампании удержания.",
        15,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 4 — cleaning and FE
    slide = blank_slide(prs)
    add_title(
        slide,
        4,
        "Предобработка, признаки и утечка",
        "Все обучаемые преобразования выполняются внутри sklearn Pipeline",
    )
    add_footer(slide, 4, "Data Scientist")
    add_panel(
        slide,
        0.62,
        1.25,
        3.85,
        1.58,
        "ПРОПУСКИ",
        (
            "TotalCharges → числовой тип; медианная импутация только на train-фолде. "
            "Пустые значения связаны с tenure = 0."
        ),
        BLUE,
        13,
    )
    add_panel(
        slide,
        0.62,
        3.02,
        3.85,
        1.58,
        "TELCO: +7 ПРИЗНАКОВ",
        (
            "tenure_group · num_addons · avg_monthly_charge · charge_growth · "
            "has_family · electronic_check · new_customer"
        ),
        CYAN,
        12,
    )
    add_panel(
        slide,
        0.62,
        4.79,
        3.85,
        1.58,
        "BANK: +5 ПРИЗНАКОВ",
        (
            "contacted_before · pdays_recency · campaign_log · season · macro_stress; "
            "pdays=999 трактуется как «никогда»."
        ),
        GREEN,
        12,
    )
    add_picture_fit(slide, "bank_duration_leakage.png", 4.75, 1.27, 7.72, 4.9)
    add_text(
        slide,
        4.92,
        6.25,
        7.35,
        0.48,
        "duration известен только после звонка: +0,148 ROC-AUC — фантомное качество, признак исключён.",
        14,
        RED,
        True,
        PP_ALIGN.CENTER,
    )

    # 5 — architecture
    slide = blank_slide(prs)
    add_title(
        slide,
        5,
        "Архитектура и MLOps-автоматизация",
        "Структурированный репозиторий и воспроизводимость одной командой",
    )
    add_footer(slide, 5, "ML Engineer")
    add_flow(
        slide,
        [
            ("download\nSHA-256", BLUE),
            ("clean + FE", CYAN),
            ("train + CV", GREEN),
            ("evaluate", ORANGE),
            ("reports + model", BLUE),
        ],
        y=1.35,
    )
    data = [
        ["Слой", "Файлы", "Назначение"],
        ["Конфигурация", "config.py", "пути, seed, реестр данных"],
        ["Данные", "download.py · datasets.py", "загрузка, checksum, очистка"],
        ["ML", "features.py · training.py", "признаки, модели, CV, порог"],
        ["Оценка", "evaluation.py · plotting.py", "метрики и графики"],
        ["Оркестрация", "pipeline.py · Makefile · dvc.yaml", "запуск одной командой"],
    ]
    add_table(slide, 0.65, 2.7, 8.0, 3.45, data, [1.55, 2.75, 3.7], 11)
    add_panel(slide, 8.92, 2.7, 3.78, 1.0, "23 ТЕСТА", "pytest: все проверки пройдены", GREEN, 13)
    add_panel(slide, 8.92, 3.88, 3.78, 1.0, "КАЧЕСТВО КОДА", "flake8 + black: 0 замечаний", BLUE, 13)
    add_panel(slide, 8.92, 5.06, 3.78, 1.0, "CI / DVC", "smoke-прогон · граф зависимостей", ORANGE, 13)
    add_text(
        slide,
        0.7,
        6.45,
        12.0,
        0.35,
        "bash scripts/run_all.sh   •   make pipeline   •   dvc repro",
        16,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 6 — methodology
    slide = blank_slide(prs)
    add_title(slide, 6, "Методология и сравнение моделей", "Честная оценка при дисбалансе и защита от утечек")
    add_footer(slide, 6, "Совместно")
    add_flow(
        slide,
        [
            ("Stratified\n80 / 20", BLUE),
            ("5-fold CV\n4 модели", CYAN),
            ("Randomized\nSearch CV", GREEN),
            ("OOF-порог\nпо F1", ORANGE),
            ("финальный\ntest", BLUE),
        ],
        y=1.28,
    )
    data = [
        ["Модель", "Семейство", "Работа с дисбалансом"],
        ["Logistic Regression", "линейная", "class_weight=balanced"],
        ["Random Forest", "бэггинг", "class_weight=balanced"],
        ["XGBoost", "бустинг", "scale_pos_weight"],
        ["LightGBM", "бустинг", "class_weight=balanced"],
    ]
    add_table(slide, 0.65, 2.68, 7.25, 3.2, data, [2.5, 1.55, 3.2], 12)
    add_panel(
        slide,
        8.22,
        2.68,
        4.48,
        1.17,
        "МЕТРИКИ",
        "Accuracy · Precision · Recall · F1 · ROC-AUC · PR-AUC",
        BLUE,
        13,
    )
    add_panel(
        slide,
        8.22,
        4.02,
        4.48,
        1.17,
        "ЗАЩИТА ОТ УТЕЧЕК",
        "preprocessing на train · duration удалён · порог по OOF",
        GREEN,
        13,
    )
    add_panel(
        slide,
        8.22,
        5.36,
        4.48,
        1.17,
        "ВОСПРОИЗВОДИМОСТЬ",
        "random_state=42 · фиксированный split · SHA-256 данных",
        ORANGE,
        13,
    )

    # 7 — Telco results
    slide = blank_slide(prs)
    add_title(
        slide,
        7,
        "Результаты: Telco Customer Churn",
        "Финальная модель — Logistic Regression, оптимальный порог 0,58",
    )
    add_footer(slide, 7, "Data Scientist")
    add_metric(slide, 0.65, 1.25, 1.82, "ROC-AUC", "0,848", BLUE, "ранжирование")
    add_metric(slide, 2.62, 1.25, 1.82, "F1", "0,637", GREEN, "+0,026 к порогу 0,5")
    add_metric(slide, 4.59, 1.25, 1.82, "Recall", "0,746", GREEN, "найдено ушедших")
    add_metric(slide, 6.56, 1.25, 1.82, "Precision", "0,556", BLUE, "точность сигналов")
    add_metric(slide, 8.53, 1.25, 1.82, "Accuracy", "0,774", BLUE, "вспомогательная")
    add_metric(slide, 10.50, 1.25, 1.82, "PR-AUC", "0,670", GREEN, "дисбаланс")
    add_picture_fit(slide, "telco_roc_curves.png", 0.75, 2.55, 6.02, 3.78)
    add_picture_fit(slide, "telco_threshold_curve.png", 6.92, 2.55, 5.7, 3.78)
    add_text(
        slide,
        0.8,
        6.45,
        11.7,
        0.38,
        "Линейная модель обошла бустинги и остаётся наиболее интерпретируемой и дешёвой в сопровождении.",
        14,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 8 — Bank results
    slide = blank_slide(prs)
    add_title(slide, 8, "Результаты: Bank Marketing", "Финальная модель — XGBoost, оптимальный порог 0,19")
    add_footer(slide, 8, "Data Scientist")
    add_metric(slide, 0.65, 1.25, 1.82, "ROC-AUC", "0,817", BLUE, "ранжирование")
    add_metric(slide, 2.62, 1.25, 1.82, "F1", "0,528", GREEN, "+0,186 к порогу 0,5")
    add_metric(slide, 4.59, 1.25, 1.82, "Recall", "0,603", GREEN, "вместо 0,227")
    add_metric(slide, 6.56, 1.25, 1.82, "Precision", "0,469", BLUE, "точность сигналов")
    add_metric(slide, 8.53, 1.25, 1.82, "Accuracy", "0,878", BLUE, "после порога")
    add_metric(slide, 10.50, 1.25, 1.82, "PR-AUC", "0,492", GREEN, "при 11,3% positive")
    add_picture_fit(slide, "bank_roc_curves.png", 0.75, 2.55, 6.02, 3.78)
    add_picture_fit(slide, "bank_threshold_curve.png", 6.92, 2.55, 5.7, 3.78)
    add_text(
        slide,
        0.8,
        6.45,
        11.7,
        0.38,
        "Порог важен не меньше модели: F1 вырос с 0,342 до 0,528, Recall — с 0,227 до 0,603.",
        14,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 9 — interpretation
    slide = blank_slide(prs)
    add_title(
        slide,
        9,
        "Интерпретация и практическая польза",
        "Permutation importance: падение ROC-AUC при перемешивании признака",
    )
    add_footer(slide, 9, "Data Scientist")
    add_picture_fit(slide, "telco_feature_importance.png", 0.55, 1.2, 6.1, 4.65)
    add_picture_fit(slide, "bank_feature_importance.png", 6.72, 1.2, 6.1, 4.65)
    add_panel(
        slide,
        0.7,
        5.92,
        5.85,
        0.82,
        "TELCO: КОМУ ПРЕДЛОЖИТЬ УДЕРЖАНИЕ",
        "Contract · tenure · tenure_group — управляемые клиентские признаки",
        BLUE,
        12,
    )
    add_panel(
        slide,
        6.78,
        5.92,
        5.85,
        0.82,
        "BANK: КОГДА И КОМУ ЗВОНИТЬ",
        "nr.employed · emp.var.rate · euribor3m — макроэкономический фон",
        GREEN,
        12,
    )

    # 10 — conclusions
    slide = blank_slide(prs)
    add_title(
        slide,
        10,
        "Выводы, ограничения и следующая итерация",
        "Результат готов к пилоту и дальнейшей продуктовой проверке",
    )
    add_footer(slide, 10, "Совместно")
    add_panel(
        slide,
        0.65,
        1.25,
        3.85,
        2.08,
        "ЧТО ПОЛУЧИЛИ",
        (
            "Telco: ROC-AUC 0,848 · F1 0,637\n"
            "Bank: ROC-AUC 0,817 · F1 0,528\n\n"
            "Модели ранжируют клиентов и помогают направлять ограниченный бюджет."
        ),
        GREEN,
        14,
    )
    add_panel(
        slide,
        4.74,
        1.25,
        3.85,
        2.08,
        "ЧТО РЕШИЛИ",
        (
            "Дисбаланс → F1 / PR-AUC\nУтечка duration → исключена\n"
            "Пропуски → pipeline-импутация\nДолгий запуск → fast-режим и DVC"
        ),
        BLUE,
        14,
    )
    add_panel(
        slide,
        8.83,
        1.25,
        3.85,
        2.08,
        "ОГРАНИЧЕНИЯ",
        (
            "Нет оценки бизнес-стоимости ошибок\nНет калибровки вероятностей\n"
            "Нет мониторинга дрейфа\nРезультаты требуют A/B-пилота"
        ),
        ORANGE,
        14,
    )
    add_panel(
        slide,
        0.65,
        3.65,
        3.85,
        1.7,
        "1. DEPLOYMENT",
        "FastAPI + Docker\nPOST /predict · healthcheck",
        BLUE,
        14,
    )
    add_panel(
        slide,
        4.74,
        3.65,
        3.85,
        1.7,
        "2. ЭКСПЕРИМЕНТЫ",
        "MLflow · калибровка\nпорог по стоимости ошибок",
        GREEN,
        14,
    )
    add_panel(
        slide,
        8.83,
        3.65,
        3.85,
        1.7,
        "3. МОНИТОРИНГ",
        "PSI / KS · деградация метрик\nавтоматическое переобучение",
        ORANGE,
        14,
    )
    add_text(
        slide,
        0.7,
        5.78,
        12.0,
        0.48,
        "Воспроизведение проекта:  bash scripts/run_all.sh",
        18,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide, 0.7, 6.37, 12.0, 0.42, "Спасибо! Готовы ответить на вопросы.", 21, BLUE, True, PP_ALIGN.CENTER
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"[ok] {path} ({path.stat().st_size / 1024 / 1024:.2f} MiB)")
