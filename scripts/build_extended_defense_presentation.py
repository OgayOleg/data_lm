#!/usr/bin/env python
"""Build an extended 20-slide PowerPoint deck for the project defense."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_defense_presentation import (
    BLUE,
    CYAN,
    DARK_BG,
    GREEN,
    MUTED,
    NAVY,
    ORANGE,
    RED,
    SLIDE_H,
    SLIDE_W,
    WHITE,
    add_box,
    add_flow,
    add_metric,
    add_panel,
    add_picture_fit,
    add_table,
    add_text,
    add_title,
    blank_slide,
)

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "reports" / "defense_presentation_20_slides.pptx"
TOTAL = 20


def footer(slide, index, role="Совместно"):
    add_text(slide, 0.55, 7.16, 5.5, 0.18, f"Введение в науку о данных  •  {role}", 9, MUTED)
    add_text(slide, 12.02, 7.16, 0.72, 0.18, f"{index}/{TOTAL}", 9, MUTED, True, PP_ALIGN.RIGHT)


def section_label(slide, text, x=0.65, y=1.23, w=2.5, color=BLUE):
    box = add_box(slide, x, y, w, 0.37, color, color)
    add_text(slide, x + 0.08, y + 0.06, w - 0.16, 0.2, text.upper(), 10, WHITE, True, PP_ALIGN.CENTER)
    return box


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # 1. Title
    slide = blank_slide(prs)
    slide.background.fill.fore_color.rgb = DARK_BG
    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), Inches(SLIDE_H)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CYAN
    accent.line.fill.background()
    add_text(slide, 0.82, 0.7, 7.2, 0.35, "ПРОЕКТ-ПРАКТИКУМ  •  DEVOPS-ИНЖЕНЕРИЯ", 12, CYAN, True)
    add_text(
        slide,
        0.82,
        1.25,
        11.65,
        1.55,
        "Разработка и автоматизация\nML-пайплайна бинарной классификации",
        34,
        WHITE,
        True,
    )
    add_text(
        slide,
        0.82,
        3.0,
        10.5,
        0.44,
        "Telco Customer Churn  ×  Bank Marketing",
        20,
        RGBColor(187, 210, 238),
        True,
    )
    add_panel(
        slide,
        0.82,
        4.02,
        5.65,
        1.36,
        "Data Scientist / Researcher",
        "ФИО: ____________________\nEDA · признаки · модели · интерпретация",
        CYAN,
        14,
    )
    add_panel(
        slide,
        6.78,
        4.02,
        5.65,
        1.36,
        "ML Engineer / DevOps Engineer",
        "ФИО: ____________________\nархитектура · тесты · DVC · CI · документация",
        GREEN,
        14,
    )
    add_text(
        slide, 0.82, 6.3, 7.4, 0.32, "Дисциплина: «Введение в науку о данных»", 13, RGBColor(187, 210, 238)
    )
    add_text(
        slide, 9.7, 6.3, 2.72, 0.32, "Группа: __________", 13, RGBColor(187, 210, 238), False, PP_ALIGN.RIGHT
    )
    add_text(slide, 11.9, 7.16, 0.55, 0.18, "1/20", 9, RGBColor(187, 210, 238), True, PP_ALIGN.RIGHT)

    # 2. Brief and goal
    slide = blank_slide(prs)
    add_title(
        slide,
        2,
        "Цель и требования проекта",
        "Сквозной воспроизводимый ML-процесс для бинарной классификации",
    )
    footer(slide, 2, "Совместно")
    add_panel(
        slide,
        0.65,
        1.35,
        3.8,
        1.55,
        "ЦЕЛЬ",
        "Закрепить EDA, подготовку данных, моделирование и принципы DevOps/MLOps в одном проекте.",
        BLUE,
        15,
    )
    add_panel(
        slide,
        4.76,
        1.35,
        3.8,
        1.55,
        "ЗАДАЧА",
        "Построить end-to-end pipeline на публичных данных объёмом более 5 000 строк.",
        GREEN,
        15,
    )
    add_panel(
        slide,
        8.87,
        1.35,
        3.8,
        1.55,
        "РЕЗУЛЬТАТ",
        "Воспроизводимый код, модели, метрики, тесты, графики, отчёт и презентация.",
        ORANGE,
        15,
    )
    add_flow(
        slide,
        [
            ("EDA", BLUE),
            ("Feature\nEngineering", CYAN),
            ("4 модели", GREEN),
            ("валидация", ORANGE),
            ("MLOps", BLUE),
        ],
        y=3.38,
    )
    add_panel(
        slide,
        0.95,
        4.78,
        11.45,
        1.25,
        "КРИТЕРИИ УСПЕХА",
        "Качество исследования  •  корректные метрики  •  отсутствие утечек  •  читаемый код  •  автоматизация  •  командная защита",  # noqa: E501
        BLUE,
        16,
    )
    add_text(
        slide,
        0.9,
        6.35,
        11.55,
        0.35,
        "Проект закрывает требования методического указания и дополнительно сравнивает два датасета.",
        15,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 3. Roles
    slide = blank_slide(prs)
    add_title(
        slide,
        3,
        "Распределение ролей в команде",
        "Ответственность разделена, решения и выводы обсуждались совместно",
    )
    footer(slide, 3, "Совместно")
    add_panel(
        slide,
        0.7,
        1.35,
        5.9,
        4.85,
        "DATA SCIENTIST / RESEARCHER",
        "1. Выбор и анализ датасетов\n2. EDA: баланс, пропуски, корреляции\n3. Feature Engineering\n4. Обучение четырёх моделей\n5. CV, тюнинг и выбор порога\n6. Интерпретация признаков\n7. Бизнес-выводы",  # noqa: E501
        CYAN,
        17,
    )
    add_panel(
        slide,
        6.75,
        1.35,
        5.9,
        4.85,
        "ML ENGINEER / DEVOPS ENGINEER",
        "1. Структура репозитория\n2. Виртуальное окружение и зависимости\n3. Рефакторинг в модули src/\n4. Модульные тесты\n5. flake8 и black\n6. Makefile, DVC и GitHub Actions\n7. Документация и воспроизводимость",  # noqa: E501
        GREEN,
        17,
    )
    add_text(
        slide,
        0.8,
        6.43,
        11.75,
        0.35,
        "Совместно: постановка эксперимента, анализ ошибок, финальный отчёт и защита.",
        15,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 4. Datasets
    slide = blank_slide(prs)
    add_title(slide, 4, "Две задачи и два датасета", "48 231 наблюдение и разные уровни дисбаланса")
    footer(slide, 4, "Data Scientist")
    data = [
        ["Параметр", "Telco Customer Churn", "Bank Marketing"],
        ["Бизнес-вопрос", "Уйдёт ли клиент?", "Подпишется ли на депозит?"],
        ["Источник", "Kaggle / IBM", "Kaggle / UCI"],
        ["Записей", "7 043", "41 188"],
        ["Исходных признаков", "20", "20"],
        ["Целевая переменная", "Churn = Yes/No", "y = yes/no"],
        ["Положительный класс", "26,5%", "11,3%"],
        ["Главная проблема", "TotalCharges", "duration — утечка"],
    ]
    add_table(slide, 0.7, 1.35, 11.95, 4.75, data, [2.7, 4.55, 4.7], 13)
    add_panel(
        slide,
        0.9,
        6.22,
        11.55,
        0.62,
        "ЗАЧЕМ ДВА ДАТАСЕТА?",
        "Проверяем переносимость методологии, а не подгонку решения под одну таблицу.",
        BLUE,
        13,
    )

    # 5. Imbalance
    slide = blank_slide(prs)
    add_title(slide, 5, "Баланс целевых классов", "Почему Accuracy недостаточно для оценки")
    footer(slide, 5, "Data Scientist")
    add_picture_fit(slide, "telco_target_balance.png", 0.65, 1.25, 5.95, 4.7)
    add_picture_fit(slide, "bank_target_balance.png", 6.73, 1.25, 5.95, 4.7)
    add_panel(
        slide,
        0.85,
        5.98,
        5.55,
        0.78,
        "TELCO: 26,5% POSITIVE",
        "Допустим умеренный дисбаланс; важны Recall и F1.",
        BLUE,
        12,
    )
    add_panel(
        slide,
        6.93,
        5.98,
        5.55,
        0.78,
        "BANK: 11,3% POSITIVE",
        "Наивная модель даст 88,7% Accuracy и нулевой Recall.",
        RED,
        12,
    )

    # 6. Telco EDA
    slide = blank_slide(prs)
    add_title(
        slide,
        6,
        "EDA Telco: кто находится в зоне риска",
        "Контракт, стаж и способ оплаты формируют главный сегмент оттока",
    )
    footer(slide, 6, "Data Scientist")
    add_picture_fit(slide, "telco_correlation_heatmap.png", 0.62, 1.25, 6.15, 4.75)
    add_panel(
        slide,
        7.03,
        1.35,
        5.62,
        1.02,
        "42,7%",
        "отток при Month-to-month против 2,8% при Two year",
        ORANGE,
        14,
    )
    add_panel(
        slide,
        7.03,
        2.57,
        5.62,
        1.02,
        "~50%",
        "отток при стаже 0–10 месяцев; после 4 лет — около 10%",
        BLUE,
        14,
    )
    add_panel(
        slide,
        7.03,
        3.79,
        5.62,
        1.02,
        "45,3%",
        "отток при Electronic check против ~15% у автоплатежей",
        GREEN,
        14,
    )
    add_panel(
        slide,
        7.03,
        5.01,
        5.62,
        1.02,
        "СЕГМЕНТ РИСКА",
        "новый клиент + месячный контракт + ручная оплата",
        RED,
        14,
    )
    add_text(
        slide,
        0.8,
        6.4,
        11.7,
        0.33,
        "EDA дал не только графики, но и адресную гипотезу для кампании удержания.",
        15,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 7. Bank EDA
    slide = blank_slide(prs)
    add_title(
        slide,
        7,
        "EDA Bank: история контактов и макроэкономика",
        "Положительный класс редкий, а признаки сильно коррелируют",
    )
    footer(slide, 7, "Data Scientist")
    add_picture_fit(slide, "bank_correlation_heatmap.png", 0.62, 1.25, 7.2, 4.8)
    add_panel(
        slide,
        8.08,
        1.35,
        4.55,
        1.15,
        "ИСТОРИЯ КАМПАНИЙ",
        "previous > 0: ранее контактированные клиенты подписываются чаще",
        BLUE,
        13,
    )
    add_panel(
        slide,
        8.08,
        2.72,
        4.55,
        1.15,
        "PDAYS = 999",
        "это сентинел «раньше не связывались», а не число дней",
        ORANGE,
        13,
    )
    add_panel(
        slide,
        8.08,
        4.09,
        4.55,
        1.15,
        "МАКРО-ФОН",
        "euribor3m, nr.employed и emp.var.rate взаимосвязаны",
        GREEN,
        13,
    )
    add_panel(
        slide,
        8.08,
        5.46,
        4.55,
        1.15,
        "DURATION",
        "сильный сигнал, но недоступен до завершения звонка",
        RED,
        13,
    )

    # 8. Cleaning
    slide = blank_slide(prs)
    add_title(
        slide, 8, "Очистка и стратегия предобработки", "Статистики вычисляются только на обучающей части"
    )
    footer(slide, 8, "Data Scientist")
    add_panel(
        slide,
        0.7,
        1.35,
        3.78,
        2.0,
        "ЧИСЛОВЫЕ ПРИЗНАКИ",
        "• pd.to_numeric для TotalCharges\n• медианная импутация\n• StandardScaler\n• без удаления реальных экстремумов",  # noqa: E501
        BLUE,
        15,
    )
    add_panel(
        slide,
        4.78,
        1.35,
        3.78,
        2.0,
        "КАТЕГОРИАЛЬНЫЕ",
        "• импутация наиболее частым\n• OneHotEncoder\n• unknown сохраняется как категория\n• неизвестные значения не ломают inference",  # noqa: E501
        CYAN,
        15,
    )
    add_panel(
        slide,
        8.86,
        1.35,
        3.78,
        2.0,
        "РАЗДЕЛЕНИЕ",
        "• Stratified split 80/20\n• random_state = 42\n• train/test сохраняют баланс\n• test не участвует в обучении",  # noqa: E501
        GREEN,
        15,
    )
    add_flow(
        slide,
        [
            ("raw CSV", BLUE),
            ("очистка", CYAN),
            ("split", GREEN),
            ("fit transform\nна train", ORANGE),
            ("transform\ntest", BLUE),
        ],
        y=4.05,
    )
    add_panel(
        slide,
        1.1,
        5.55,
        11.1,
        0.92,
        "ГАРАНТИЯ",
        "SimpleImputer, StandardScaler и OneHotEncoder находятся внутри sklearn Pipeline и переобучаются на каждом train-фолде.",  # noqa: E501
        GREEN,
        14,
    )

    # 9. Leakage
    slide = blank_slide(prs)
    add_title(
        slide, 9, "Эксперимент с утечкой duration", "Высокая метрика не всегда означает пригодную модель"
    )
    footer(slide, 9, "Data Scientist")
    add_picture_fit(slide, "bank_duration_leakage.png", 0.7, 1.25, 7.6, 5.25)
    add_metric(slide, 8.62, 1.55, 3.55, "Без duration", "0,802", BLUE, "честный ROC-AUC CV")
    add_metric(slide, 8.62, 2.92, 3.55, "С duration", "0,950", RED, "фантомный ROC-AUC CV")
    add_metric(slide, 8.62, 4.29, 3.55, "Разница", "+0,148", ORANGE, "искусственное улучшение")
    add_panel(slide, 8.55, 5.65, 3.7, 0.88, "РЕШЕНИЕ", "duration исключён до обучения", GREEN, 13)

    # 10. Feature engineering
    slide = blank_slide(prs)
    add_title(slide, 10, "Feature Engineering", "Создаём признаки, доступные в момент реального предсказания")
    footer(slide, 10, "Data Scientist")
    add_panel(
        slide,
        0.7,
        1.35,
        5.92,
        4.8,
        "TELCO: 7 НОВЫХ ПРИЗНАКОВ",
        "num_addons — количество услуг\ntenure_group — сегмент стажа\navg_monthly_charge — средняя оплата\ncharge_growth — изменение стоимости\nhas_family — семейный статус\nis_electronic_check — ручная оплата\nis_new_customer — новый клиент",  # noqa: E501
        BLUE,
        16,
    )
    add_panel(
        slide,
        6.75,
        1.35,
        5.92,
        4.8,
        "BANK: 5 НОВЫХ ПРИЗНАКОВ",
        "was_contacted_before — история контактов\npdays_recency — давность / never\ncampaign_log — логарифм числа звонков\nseason — сезон контакта\nmacro_stress — агрегат макропоказателей\n\nduration удалён как утечка",  # noqa: E501
        GREEN,
        16,
    )
    add_text(
        slide,
        0.8,
        6.42,
        11.75,
        0.34,
        "Созданные признаки вошли в топ permutation importance и подтвердили полезность этапа.",
        15,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 11. No leakage pipeline
    slide = blank_slide(prs)
    add_title(
        slide,
        11,
        "Как предотвращается утечка данных",
        "Pipeline объединяет preprocessing и модель в один обучаемый объект",
    )
    footer(slide, 11, "ML Engineer")
    add_flow(
        slide,
        [
            ("train fold", BLUE),
            ("imputer", CYAN),
            ("scaler +\none-hot", GREEN),
            ("model.fit", ORANGE),
            ("validation\nfold", BLUE),
        ],
        y=1.45,
    )
    add_panel(
        slide,
        0.75,
        3.05,
        3.75,
        2.25,
        "РИСК 1: PREPROCESSING",
        "Нельзя вычислять медиану, среднее или категории на всей таблице. Решение: ColumnTransformer внутри Pipeline.",  # noqa: E501
        RED,
        15,
    )
    add_panel(
        slide,
        4.8,
        3.05,
        3.75,
        2.25,
        "РИСК 2: TARGET LEAKAGE",
        "duration удалён до формирования модельной матрицы. Эффект утечки измерен отдельным экспериментом.",
        ORANGE,
        15,
    )
    add_panel(
        slide,
        8.85,
        3.05,
        3.75,
        2.25,
        "РИСК 3: ПОРОГ",
        "Оптимальный threshold выбирается по out-of-fold прогнозам train, а не по тестовой выборке.",
        GREEN,
        15,
    )
    add_panel(
        slide,
        1.25,
        5.75,
        10.85,
        0.8,
        "АВТОМАТИЧЕСКАЯ ПРОВЕРКА",
        "Тест test_preprocessor_fitted_on_train_only фиксирует основную гарантию в коде.",
        BLUE,
        14,
    )

    # 12. Repository architecture
    slide = blank_slide(prs)
    add_title(
        slide,
        12,
        "Архитектура кодовой базы",
        "Исследовательская логика вынесена из ноутбуков в повторно используемые модули",
    )
    footer(slide, 12, "ML Engineer")
    data = [
        ["Слой", "Файлы", "Ответственность"],
        ["Конфигурация", "src/config.py", "пути, seed, реестр данных, SHA-256"],
        ["Данные", "download.py · datasets.py", "загрузка, проверка, очистка"],
        ["Признаки", "features.py · preprocessing.py", "FE и ColumnTransformer"],
        ["Модели", "models.py · training.py", "4 модели, CV, tuning, threshold"],
        ["Оценка", "evaluation.py · plotting.py", "метрики и 19 графиков"],
        ["Оркестрация", "pipeline.py", "единый end-to-end запуск"],
        ["Интерфейсы", "Makefile · scripts/ · dvc.yaml", "автоматизация и воспроизведение"],
    ]
    add_table(slide, 0.7, 1.35, 11.95, 4.95, data, [2.15, 3.3, 6.5], 12)
    add_text(
        slide,
        0.85,
        6.48,
        11.65,
        0.32,
        "Ноутбуки импортируют функции из src/, а не дублируют производственный код.",
        15,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 13. Automation
    slide = blank_slide(prs)
    add_title(
        slide,
        13,
        "Автоматизация и качество кода",
        "От клонирования репозитория до метрик и модели — одна команда",
    )
    footer(slide, 13, "ML Engineer")
    add_metric(slide, 0.75, 1.35, 2.65, "Unit tests", "23", GREEN, "pytest: пройдены")
    add_metric(slide, 3.62, 1.35, 2.65, "Линтинг", "0", BLUE, "flake8 замечаний")
    add_metric(slide, 6.49, 1.35, 2.65, "Модели", "4", GREEN, "единая политика")
    add_metric(slide, 9.36, 1.35, 2.65, "Графики", "19", BLUE, "автогенерация")
    add_panel(
        slide,
        0.75,
        2.8,
        3.72,
        2.45,
        "MAKE / SCRIPTS",
        "make setup\nmake data\nmake pipeline\nmake test\nmake check\n\nbash scripts/run_all.sh",
        BLUE,
        16,
    )
    add_panel(
        slide,
        4.8,
        2.8,
        3.72,
        2.45,
        "DVC",
        "download → train_telco\n             ↘ train_bank\n\nКэширование стадий, зависимости и параметры эксперимента.",  # noqa: E501
        GREEN,
        16,
    )
    add_panel(
        slide,
        8.85,
        2.8,
        3.72,
        2.45,
        "GITHUB ACTIONS",
        "Python 3.11 / 3.12\nflake8 + black\npytest + coverage\nTelco smoke pipeline\nартефакты CI",
        ORANGE,
        16,
    )
    add_text(
        slide,
        0.8,
        5.7,
        11.7,
        0.58,
        "Быстрый режим: --skip-tuning --skip-leakage-experiment. Полный режим создаёт модели, CSV, JSON и PNG.",  # noqa: E501
        15,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    # 14. Methodology/models
    slide = blank_slide(prs)
    add_title(
        slide,
        14,
        "Методология моделирования",
        "Стратификация, кросс-валидация, тюнинг и выбор operating point",
    )
    footer(slide, 14, "Совместно")
    add_flow(
        slide,
        [
            ("split\n80 / 20", BLUE),
            ("5-fold CV", CYAN),
            ("4 модели", GREEN),
            ("Randomized\nSearch", ORANGE),
            ("OOF-порог", BLUE),
        ],
        y=1.25,
    )
    data = [
        ["Модель", "Семейство", "Дисбаланс", "Роль"],
        ["Logistic Regression", "линейная", "class_weight", "интерпретируемый baseline"],
        ["Random Forest", "бэггинг", "class_weight", "нелинейный benchmark"],
        ["XGBoost", "бустинг", "scale_pos_weight", "точные нелинейные зависимости"],
        ["LightGBM", "бустинг", "class_weight", "быстрый gradient boosting"],
    ]
    add_table(slide, 0.7, 2.65, 11.95, 3.15, data, [2.65, 2.05, 2.25, 5.0], 12)
    add_panel(
        slide,
        0.9,
        6.05,
        11.55,
        0.7,
        "МЕТРИКИ",
        "Accuracy · Precision · Recall · F1 · ROC-AUC · PR-AUC. Основной ranking: ROC-AUC; порог: максимум OOF F1.",  # noqa: E501
        BLUE,
        13,
    )

    # 15. CV stability
    slide = blank_slide(prs)
    add_title(
        slide,
        15,
        "Стабильность на кросс-валидации",
        "Сравниваем не одну случайную выборку, а распределение качества по фолдам",
    )
    footer(slide, 15, "Data Scientist")
    add_picture_fit(slide, "telco_cv_roc_auc_boxplot.png", 0.62, 1.25, 6.0, 4.95)
    add_picture_fit(slide, "bank_cv_roc_auc_boxplot.png", 6.72, 1.25, 6.0, 4.95)
    add_panel(
        slide,
        1.0,
        6.13,
        5.25,
        0.65,
        "TELCO",
        "Logistic Regression стабильна и конкурентна бустингам.",
        BLUE,
        12,
    )
    add_panel(
        slide,
        7.1,
        6.13,
        5.25,
        0.65,
        "BANK",
        "XGBoost и LightGBM дают близкие распределения ROC-AUC.",
        GREEN,
        12,
    )

    # 16. Telco results
    slide = blank_slide(prs)
    add_title(slide, 16, "Финальный результат Telco", "Logistic Regression, threshold = 0,58")
    footer(slide, 16, "Data Scientist")
    add_metric(slide, 0.65, 1.25, 1.82, "ROC-AUC", "0,848", BLUE)
    add_metric(slide, 2.62, 1.25, 1.82, "F1", "0,637", GREEN)
    add_metric(slide, 4.59, 1.25, 1.82, "Recall", "0,746", GREEN)
    add_metric(slide, 6.56, 1.25, 1.82, "Precision", "0,556", BLUE)
    add_metric(slide, 8.53, 1.25, 1.82, "Accuracy", "0,774", BLUE)
    add_metric(slide, 10.5, 1.25, 1.82, "PR-AUC", "0,670", GREEN)
    add_picture_fit(slide, "telco_roc_curves.png", 0.65, 2.55, 6.1, 3.75)
    add_picture_fit(slide, "telco_confusion_matrices.png", 6.88, 2.55, 5.8, 3.75)
    add_text(
        slide,
        0.8,
        6.42,
        11.7,
        0.34,
        "Модель находит около 75% реально уходящих клиентов и проще бустингов в сопровождении.",
        15,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 17. Bank results
    slide = blank_slide(prs)
    add_title(slide, 17, "Финальный результат Bank", "XGBoost, threshold = 0,19")
    footer(slide, 17, "Data Scientist")
    add_metric(slide, 0.65, 1.25, 1.82, "ROC-AUC", "0,817", BLUE)
    add_metric(slide, 2.62, 1.25, 1.82, "F1", "0,528", GREEN)
    add_metric(slide, 4.59, 1.25, 1.82, "Recall", "0,603", GREEN)
    add_metric(slide, 6.56, 1.25, 1.82, "Precision", "0,469", BLUE)
    add_metric(slide, 8.53, 1.25, 1.82, "Accuracy", "0,878", BLUE)
    add_metric(slide, 10.5, 1.25, 1.82, "PR-AUC", "0,492", GREEN)
    add_picture_fit(slide, "bank_roc_curves.png", 0.65, 2.55, 6.1, 3.75)
    add_picture_fit(slide, "bank_confusion_matrices.png", 6.88, 2.55, 5.8, 3.75)
    add_text(
        slide,
        0.8,
        6.42,
        11.7,
        0.34,
        "При редком positive-классе модель обнаруживает около 60% потенциальных подписчиков.",
        15,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )

    # 18. Threshold
    slide = blank_slide(prs)
    add_title(
        slide,
        18,
        "Порог решения — часть бизнес-модели",
        "Стандартный threshold 0,5 не оптимален при дисбалансе",
    )
    footer(slide, 18, "Data Scientist")
    add_picture_fit(slide, "telco_threshold_curve.png", 0.62, 1.25, 6.0, 4.8)
    add_picture_fit(slide, "bank_threshold_curve.png", 6.72, 1.25, 6.0, 4.8)
    add_panel(slide, 0.95, 5.95, 5.35, 0.82, "TELCO", "threshold 0,58: F1 0,611 → 0,637", BLUE, 13)
    add_panel(
        slide, 7.05, 5.95, 5.35, 0.82, "BANK", "threshold 0,19: F1 0,342 → 0,528; Recall ×2,7", GREEN, 13
    )

    # 19. Importance and business
    slide = blank_slide(prs)
    add_title(slide, 19, "Интерпретация и практическая польза", "Permutation importance по падению ROC-AUC")
    footer(slide, 19, "Data Scientist")
    add_picture_fit(slide, "telco_feature_importance.png", 0.55, 1.18, 6.1, 4.65)
    add_picture_fit(slide, "bank_feature_importance.png", 6.72, 1.18, 6.1, 4.65)
    add_panel(
        slide,
        0.7,
        5.9,
        5.85,
        0.82,
        "TELCO: КОМУ ПРЕДЛОЖИТЬ УДЕРЖАНИЕ",
        "Contract · tenure · tenure_group — управляемые клиентские признаки",
        BLUE,
        12,
    )
    add_panel(
        slide,
        6.78,
        5.9,
        5.85,
        0.82,
        "BANK: КОГДА И КОМУ ЗВОНИТЬ",
        "nr.employed · emp.var.rate · euribor3m — макроэкономический фон",
        GREEN,
        12,
    )

    # 20. Conclusions
    slide = blank_slide(prs)
    add_title(
        slide, 20, "Выводы и следующая итерация", "Воспроизводимый результат готов к пилотной бизнес-проверке"
    )
    footer(slide, 20, "Совместно")
    add_panel(
        slide,
        0.65,
        1.28,
        3.82,
        1.78,
        "РЕЗУЛЬТАТ",
        "Telco: ROC-AUC 0,848 · F1 0,637\nBank: ROC-AUC 0,817 · F1 0,528\n\nПайплайн воспроизводится одной командой.",  # noqa: E501
        GREEN,
        14,
    )
    add_panel(
        slide,
        4.75,
        1.28,
        3.82,
        1.78,
        "ТРУДНОСТИ",
        "Дисбаланс классов\nПропуски и сентинелы\nУтечка duration\nДолгие CPU-прогоны",
        ORANGE,
        14,
    )
    add_panel(
        slide,
        8.85,
        1.28,
        3.82,
        1.78,
        "ОГРАНИЧЕНИЯ",
        "Нет денежной стоимости ошибок\nНет калибровки\nНет мониторинга дрейфа\nНужен A/B-пилот",
        RED,
        14,
    )
    add_panel(
        slide, 0.65, 3.43, 3.82, 1.52, "1. DEPLOYMENT", "FastAPI + Docker\nверсионирование модели", BLUE, 14
    )
    add_panel(
        slide,
        4.75,
        3.43,
        3.82,
        1.52,
        "2. ЭКСПЕРИМЕНТЫ",
        "MLflow · SHAP\nпорог по стоимости ошибок",
        GREEN,
        14,
    )
    add_panel(
        slide,
        8.85,
        3.43,
        3.82,
        1.52,
        "3. МОНИТОРИНГ",
        "PSI / KS · качество\nавтоматическое переобучение",
        ORANGE,
        14,
    )
    add_text(slide, 0.8, 5.45, 11.7, 0.4, "bash scripts/run_all.sh", 18, NAVY, True, PP_ALIGN.CENTER)
    add_text(
        slide, 0.8, 6.12, 11.7, 0.52, "Спасибо! Готовы ответить на вопросы.", 23, BLUE, True, PP_ALIGN.CENTER
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"[ok] {path} ({path.stat().st_size / 1024 / 1024:.2f} MiB)")
